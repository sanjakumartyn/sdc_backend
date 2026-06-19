import os
import re
import json
import datetime
import requests
import logging
from typing import Any, Dict, List, Tuple
from django.conf import settings

# Libraries for document generation
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from common.utils.gemini import call_gemini_chat
from features.companyAnalysis.service import CompanyAnalysisService
from features.companydata.service import CompanyDataService

logger = logging.getLogger(__name__)

class DocumentGenerationService:
    @staticmethod
    def generate_document(company: str, document_type: str) -> Dict[str, Any]:
        """
        Retrieves internal and external data, synthesizes it with an LLM,
        and generates downloadable documents in the configured formats.
        """
        # Ensure media directories exist
        doc_dir = os.path.join(settings.MEDIA_ROOT, "documents")
        os.makedirs(doc_dir, exist_ok=True)
        
        # 1. Fetch context data from MongoDB and external sources
        context = DocumentGenerationService._build_enriched_context(company)
        
        # 2. Call LLM to synthesize report contents
        sections = DocumentGenerationService._synthesize_document_content(company, document_type, context)
        
        # 3. Generate requested files
        company_safe = re.sub(r'[^a-zA-Z0-9_]', '_', company)
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        
        downloads = {}
        
        # Determine formats to generate based on documentType
        if document_type == "excel_sheet":
            # Only Excel
            xlsx_filename = f"opportunity_sheet_{company_safe}_{timestamp}.xlsx"
            xlsx_path = os.path.join(doc_dir, xlsx_filename)
            DocumentGenerationService._generate_excel_file(company, context, xlsx_path)
            downloads["xlsx"] = f"{settings.MEDIA_URL}documents/{xlsx_filename}"
        elif document_type == "opportunity_report":
            # DOCX, PDF, and Excel
            docx_filename = f"opportunity_report_{company_safe}_{timestamp}.docx"
            pdf_filename = f"opportunity_report_{company_safe}_{timestamp}.pdf"
            xlsx_filename = f"opportunity_sheet_{company_safe}_{timestamp}.xlsx"
            
            docx_path = os.path.join(doc_dir, docx_filename)
            pdf_path = os.path.join(doc_dir, pdf_filename)
            xlsx_path = os.path.join(doc_dir, xlsx_filename)
            
            DocumentGenerationService._generate_docx_file("Opportunity Report", company, sections, docx_path)
            DocumentGenerationService._generate_pdf_file("Opportunity Report", company, sections, pdf_path)
            DocumentGenerationService._generate_excel_file(company, context, xlsx_path)
            
            downloads["docx"] = f"{settings.MEDIA_URL}documents/{docx_filename}"
            downloads["pdf"] = f"{settings.MEDIA_URL}documents/{pdf_filename}"
            downloads["xlsx"] = f"{settings.MEDIA_URL}documents/{xlsx_filename}"
        else:
            # DOCX and PDF (proposal, meeting_brief, qbr)
            doc_title = document_type.replace("_", " ").title()
            docx_filename = f"{document_type}_{company_safe}_{timestamp}.docx"
            pdf_filename = f"{document_type}_{company_safe}_{timestamp}.pdf"
            
            docx_path = os.path.join(doc_dir, docx_filename)
            pdf_path = os.path.join(doc_dir, pdf_filename)
            
            DocumentGenerationService._generate_docx_file(doc_title, company, sections, docx_path)
            DocumentGenerationService._generate_pdf_file(doc_title, company, sections, pdf_path)
            
            downloads["docx"] = f"{settings.MEDIA_URL}documents/{docx_filename}"
            downloads["pdf"] = f"{settings.MEDIA_URL}documents/{pdf_filename}"
            
        return {
            "status": "success",
            "message": f"{document_type.replace('_', ' ').title()} Generated Successfully",
            "downloads": downloads,
            "timestamp": datetime.datetime.now().isoformat()
        }
        
    @staticmethod
    def _build_enriched_context(company: str) -> Dict[str, Any]:
        """Queries local MongoDB data and calls upstream RAG/intelligence services to construct context."""
        context = {}
        
        # A. Query local MongoDB collections for matches
        try:
            db = CompanyDataService._get_db()
            regex = re.compile(company, re.IGNORECASE)
            
            context["crm_records"] = list(db['CRM Records'].find({"companyName": regex}, {"_id": 0}))
            context["opportunity_history"] = list(db['Opportunity History'].find({"companyName": regex}, {"_id": 0}))
            context["meeting_notes"] = list(db['past sales and meeting records'].find({"companyName": regex}, {"_id": 0}))
            context["proposals"] = list(db['Proposal documents'].find({"companyName": regex}, {"_id": 0}))
            
            # Grab some product and case study details for RAG context fallback
            context["catalog_products"] = list(db['product details'].find({}, {"_id": 0}).limit(10))
            context["catalog_case_studies"] = list(db['case studies'].find({}, {"_id": 0}).limit(10))
        except Exception as e:
            logger.error(f"Error reading MongoDB for context: {e}", exc_info=True)
            context["crm_records"] = []
            context["opportunity_history"] = []
            context["meeting_notes"] = []
            context["proposals"] = []
            context["catalog_products"] = []
            context["catalog_case_studies"] = []
            
        # B. Grab company analysis RAG context using the existing company analysis service
        try:
            analysis_payload = {"company_name": company}
            analysis_context = CompanyAnalysisService._build_context(analysis_payload, uploaded_files=[])
            context["analysis_context"] = analysis_context
        except Exception as e:
            logger.error(f"Error calling CompanyAnalysisService._build_context: {e}", exc_info=True)
            context["analysis_context"] = {}
            
        return context

    @staticmethod
    def _synthesize_document_content(company: str, document_type: str, context: Dict[str, Any]) -> Dict[str, str]:
        """Sends context to Gemini (or Groq) to synthesize detailed, clean sections in Markdown."""
        # 1. Define required keys based on documentType
        sections_map = {
            "proposal": [
                "Executive Summary",
                "Customer Challenges",
                "Recommended Solutions",
                "Product Mapping",
                "Relevant Case Studies",
                "Estimated Business Value",
                "ROI Impact",
                "Implementation Approach",
                "Next Steps"
            ],
            "meeting_brief": [
                "Company Overview",
                "Stakeholder Analysis",
                "Business Priorities",
                "Pain Points",
                "Discussion Topics",
                "Potential Objections",
                "Recommended Talking Points",
                "Meeting Objectives"
            ],
            "opportunity_report": [
                "Strategic Fit Score",
                "Opportunity Score",
                "Estimated Deal Value",
                "Recommended Products",
                "Product Match Analysis",
                "Risk Assessment",
                "Competitive Position",
                "Win Probability",
                "Recommended Actions"
            ],
            "qbr": [
                "Account Summary",
                "Pipeline Overview",
                "Revenue Potential",
                "Customer Engagement Status",
                "Opportunity Progress",
                "Recommendations",
                "Next Quarter Strategy"
            ]
        }
        
        required_keys = sections_map.get(document_type, [])
        if not required_keys:
            return {}
            
        # Prepare context payload string (safely serialized)
        try:
            context_str = json.dumps(context, ensure_ascii=True, default=str)
        except Exception:
            context_str = str(context)
            
        if len(context_str) > 15000:
            context_str = context_str[:15000] + "\n... (context truncated)"
            
        system_prompt = (
            "You are an expert enterprise sales copilot for GrowthlensAI.\n"
            "Your task is to generate the textual content for a highly professional business document.\n"
            "You MUST structure your output in clean Markdown using Level 1 headings (# Heading Name) for EACH of the following sections:\n"
            + "\n".join([f"# {key}" for key in required_keys]) + "\n\n"
            "Rules:\n"
            "1. Output markdown headings exactly matching the names listed above.\n"
            "2. Under each heading, write detailed, multi-paragraph, professional enterprise content. Do not use generic placeholders or incomplete templates.\n"
            "3. Ground all findings, estimates, and product details in the provided context. Incorporate details from CRM history, RAG matches, meeting pain points, and OCR documents if present.\n"
            "4. If certain details are missing, construct highly plausible, professional industry-specific recommendations for GrowthlensAI products.\n"
        )
        
        user_message = (
            f"Generate a professional '{document_type.replace('_', ' ').title()}' for the target company '{company}'.\n\n"
            f"Here is the context data:\n{context_str}"
        )
        
        # 2. Call LLM
        response_text = DocumentGenerationService._call_llm(system_prompt, user_message)
        
        # 3. Parse Markdown response
        try:
            # Try level 1 headings first
            pattern = re.compile(r'^#\s+(.*?)\s*\n(.*?)(?=(?:^#\s+|\Z))', re.MULTILINE | re.DOTALL)
            matches = pattern.findall(response_text)
            
            # If no level 1 headings, try level 2 headings
            if not matches:
                pattern = re.compile(r'^##\s+(.*?)\s*\n(.*?)(?=(?:^##\s+|\Z))', re.MULTILINE | re.DOTALL)
                matches = pattern.findall(response_text)
                
            # Build a raw dict of normalized heading -> content
            raw_sections = {}
            for heading, content in matches:
                normalized_heading = heading.strip().lower().replace("#", "").replace(":", "").replace("*", "").replace("_", "").strip()
                raw_sections[normalized_heading] = content.strip()
                
            # Match them back to required keys
            final_sections = {}
            for key in required_keys:
                normalized_key = key.strip().lower().replace(":", "").replace("*", "").replace("_", "").strip()
                if normalized_key in raw_sections:
                    final_sections[key] = raw_sections[normalized_key]
                else:
                    # Try finding key as substring in raw headings or vice-versa
                    matched_content = None
                    for raw_h, raw_c in raw_sections.items():
                        if normalized_key in raw_h or raw_h in normalized_key:
                            matched_content = raw_c
                            break
                    if matched_content is not None:
                        final_sections[key] = matched_content
                    else:
                        final_sections[key] = (
                            f"Detailed analysis and recommendations regarding {key.lower()} for {company} "
                            "based on sales CRM records and GrowthlensAI solutions catalog."
                        )
            
            # Verify if we got at least one actual parsed section with substantial content.
            # If everything was completely empty or failed to parse, fall back.
            total_content_len = sum(len(v) for v in final_sections.values())
            fallback_count = sum(1 for v in final_sections.values() if "based on sales CRM records and GrowthlensAI solutions catalog." in v)
            
            if fallback_count == len(required_keys) or total_content_len < 100:
                raise ValueError("Parsed content is empty or contains only fallback text.")
                
            return final_sections
        except Exception as e:
            logger.error(f"Failed to parse LLM markdown response for document content: {e}. Raw response: {response_text}", exc_info=True)
            # Safe fallback text-based dictionary
            fallback = {}
            for key in required_keys:
                fallback[key] = f"Detailed analysis and recommendations regarding {key.lower()} for {company} based on sales CRM records and GrowthlensAI solutions catalog."
            return fallback

    @staticmethod
    def _call_llm(system_prompt: str, user_message: str) -> str:
        """Helper to invoke Gemini with a fallback to Groq completions."""
        if os.getenv("GEMINI_API_KEY", "").strip():
            try:
                result = call_gemini_chat(
                    messages=[
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": user_message}
                    ],
                    temperature=0.2,
                    timeout=float(os.getenv("GEMINI_TIMEOUT", "45")),
                )
                content = (result.get("content") or "").strip()
                if content:
                    return content
            except Exception as e:
                logger.warning(f"Gemini document synthesis failed: {e}. Retrying with Groq...")

        groq_api_key = os.getenv("GROQ_API_KEY", "").strip()
        if groq_api_key:
            groq_url = os.getenv("GROQ_API_URL", "https://api.groq.com/openai/v1/chat/completions").strip()
            groq_model = os.getenv("GROQ_MODEL", "llama-3.1-8b-instant").strip()
            headers = {
                "Authorization": f"Bearer {groq_api_key}",
                "Content-Type": "application/json"
            }
            body = {
                "model": groq_model,
                "messages": [
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_message}
                ],
                "temperature": 0.2
            }
            try:
                response = requests.post(groq_url, headers=headers, json=body, timeout=30)
                response.raise_for_status()
                choices = response.json().get("choices", [])
                if choices:
                    content = choices[0].get("message", {}).get("content", "").strip()
                    if content:
                        return content
            except Exception as e:
                logger.error(f"Groq document synthesis fallback failed: {e}")
                
        # If both fail, raise service unavailable or return static mock
        return "{}"

    @staticmethod
    def _generate_docx_file(doc_title: str, company: str, sections: Dict[str, str], file_path: str):
        """Generates a professional DOCX document using python-docx."""
        doc = Document()
        
        # Configure standard 1-inch margins
        for s in doc.sections:
            s.top_margin = Inches(1)
            s.bottom_margin = Inches(1)
            s.left_margin = Inches(1)
            s.right_margin = Inches(1)
            
        # Design system styles
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Segoe UI'
        font.size = Pt(11)
        font.color.rgb = RGBColor(51, 65, 85) # Slate 700
        
        # Document Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run(f"ENTERPRISE {doc_title.upper()}")
        run_title.font.name = 'Segoe UI'
        run_title.font.size = Pt(24)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
        
        # Metadata / Subtitle
        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run(
            f"Prepared for: {company}\n"
            f"Date: {datetime.date.today().strftime('%B %d, %Y')}\n"
            f"Generated via: SalesIntelAI Copilot Dashboard"
        )
        run_sub.font.size = Pt(10)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(100, 116, 139) # Slate 500
        
        # Section separator line
        doc.add_paragraph("-" * 75)
        
        # Add content sections
        for title, text in sections.items():
            # Section Heading
            p_h = doc.add_heading(level=1)
            p_h.paragraph_format.space_before = Pt(16)
            p_h.paragraph_format.space_after = Pt(6)
            p_h.paragraph_format.keep_with_next = True
            
            run_h = p_h.add_run(title)
            run_h.font.name = 'Segoe UI'
            run_h.font.size = Pt(14)
            run_h.font.bold = True
            run_h.font.color.rgb = RGBColor(37, 99, 235) # Royal Blue 600
            
            # Content Paragraphs
            lines = text.split("\n")
            for line in lines:
                line_str = line.strip()
                if not line_str:
                    continue
                
                # Check for bullet points
                if line_str.startswith("-") or line_str.startswith("*"):
                    p_item = doc.add_paragraph(style='List Bullet')
                    p_item.paragraph_format.space_after = Pt(4)
                    p_item.add_run(line_str.lstrip("-*").strip())
                else:
                    p_p = doc.add_paragraph()
                    p_p.paragraph_format.space_after = Pt(8)
                    p_p.paragraph_format.line_spacing = 1.15
                    p_p.add_run(line_str)
                    
        doc.save(file_path)

    @staticmethod
    def _generate_pdf_file(doc_title: str, company: str, sections: Dict[str, str], file_path: str):
        """Generates a professional PDF document using reportlab."""
        doc = SimpleDocTemplate(
            file_path,
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54
        )
        styles = getSampleStyleSheet()
        
        # Custom Paragraph styles
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#0f172a'),
            spaceAfter=10
        )
        
        subtitle_style = ParagraphStyle(
            'DocSub',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=9,
            leading=13,
            textColor=colors.HexColor('#64748b'),
            spaceAfter=16
        )
        
        heading_style = ParagraphStyle(
            'DocHeading',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=17,
            textColor=colors.HexColor('#2563eb'),
            spaceBefore=12,
            spaceAfter=6,
            keepWithNext=True
        )
        
        body_style = ParagraphStyle(
            'DocBody',
            parent=styles['BodyText'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13.5,
            textColor=colors.HexColor('#334155'),
            spaceAfter=6
        )
        
        bullet_style = ParagraphStyle(
            'DocBullet',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13.5,
            textColor=colors.HexColor('#334155'),
            leftIndent=15,
            firstLineIndent=-8,
            spaceAfter=4
        )
        
        story = []
        
        # Add Title
        story.append(Paragraph(f"ENTERPRISE {doc_title.upper()}", title_style))
        
        # Add Subtitle
        sub_text = (
            f"Prepared for: {company}<br/>"
            f"Date: {datetime.date.today().strftime('%B %d, %Y')}<br/>"
            f"Generated via: SalesIntelAI Copilot Dashboard"
        )
        story.append(Paragraph(sub_text, subtitle_style))
        story.append(Spacer(1, 8))
        
        # Divider Line
        sep_table = Table([['']], colWidths=[letter[0] - 108])
        sep_table.setStyle(TableStyle([
            ('LINEABOVE', (0,0), (-1,-1), 1.2, colors.HexColor('#cbd5e1')),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ]))
        story.append(sep_table)
        story.append(Spacer(1, 12))
        
        # Build document pages
        for title, text in sections.items():
            story.append(Paragraph(title, heading_style))
            lines = text.split("\n")
            for line in lines:
                line_str = line.strip()
                if not line_str:
                    continue
                if line_str.startswith("-") or line_str.startswith("*"):
                    bullet_content = f"&bull; {line_str.lstrip('-*').strip()}"
                    story.append(Paragraph(bullet_content, bullet_style))
                else:
                    story.append(Paragraph(line_str, body_style))
            story.append(Spacer(1, 8))
            
        doc.build(story)

    @staticmethod
    def _generate_excel_file(company: str, context: Dict[str, Any], file_path: str):
        """Generates a premium styled Excel spreadsheet containing active deals & opportunities."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Opportunities"
        ws.views.sheetView[0].showGridLines = True
        
        # 1. Gather opportunities data from context
        opp_records = context.get("opportunity_history", [])
        crm_records = context.get("crm_records", [])
        
        data_rows = []
        
        if opp_records:
            # Use actual Mongo records
            for opp in opp_records:
                # Opportunity Name, Company, Recommended Product, Match Score, Estimated Deal Value, Priority, Status, Next Action
                # If fields aren't exactly matching, provide safe fallback
                opp_name = opp.get("opportunityId") or opp.get("proposalId") or "Industrial Upgrade"
                comp = opp.get("companyName") or company
                products = opp.get("competitors") # Or placeholder / productsRecommended
                prod_str = ", ".join(products) if isinstance(products, list) else "EcoShield Bio-Coating A1"
                match_score = opp.get("winProbability") or 85
                deal_val = opp.get("dealValue") or 1200000
                priority = "High" if deal_val > 1000000 else "Medium"
                status = opp.get("opportunityStatus") or "Open"
                next_action = opp.get("nextAction") or "Deliver compliance documentation"
                
                data_rows.append([opp_name, comp, prod_str, match_score, deal_val, priority, status, next_action])
        
        # If no explicit opportunity records, fallback and synthesize some rows using CRM deals / RAG recommendations
        if not data_rows:
            # Let's synthesize based on CRM or fallback
            crm_deal_val = 2100000
            crm_stage = "Proposal"
            crm_notes = "Interested in EcoShield Bio-Coating A1."
            if crm_records:
                first = crm_records[0]
                crm_deal_val = first.get("estimatedDealValue") or 2100000
                crm_stage = first.get("dealStage") or "Proposal"
                crm_notes = first.get("notes") or "Interested in ESG coating materials."
                
            data_rows = [
                [
                    "VOC Reduction & Coating Upgrade",
                    company,
                    "EcoShield Bio-Coating A1",
                    89,
                    crm_deal_val,
                    "High" if crm_deal_val > 1000000 else "Medium",
                    crm_stage,
                    "Deliver ESG compliance report"
                ],
                [
                    "Factory Equipment IoT Integration",
                    company,
                    "NovaSense Tracker",
                    74,
                    450000,
                    "Medium",
                    "Discovery",
                    "Arrange product demo and pilot"
                ]
            ]
            
        # 2. Build Excel Structure & Styles
        headers = [
            "Opportunity Name",
            "Company",
            "Recommended Product",
            "Match Score",
            "Estimated Deal Value",
            "Priority",
            "Status",
            "Next Action"
        ]
        
        # Merge title row
        ws.merge_cells("A1:H1")
        title = ws["A1"]
        title.value = f"Enterprise Opportunity Sheet - {company}"
        title.font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
        title.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Navy
        title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36
        
        # Subtitle row
        ws.merge_cells("A2:H2")
        sub = ws["A2"]
        sub.value = f"Generated by SalesIntelAI Copilot | Date: {datetime.date.today().strftime('%B %d, %Y')}"
        sub.font = Font(name="Segoe UI", size=9.5, italic=True, color="475569")
        sub.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 18
        
        # Empty separation row
        ws.row_dimensions[3].height = 8
        
        # Table Header Row
        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid") # Royal Blue
        header_font = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
        border_thin = Side(style="thin", color="CBD5E1")
        border_medium = Side(style="medium", color="1E293B")
        header_border = Border(top=border_thin, bottom=border_medium, left=border_thin, right=border_thin)
        
        for col_idx, header_val in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.value = header_val
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = header_border
            
        ws.row_dimensions[4].height = 26
        
        # Data Rows Formatting
        data_border = Border(top=border_thin, bottom=border_thin, left=border_thin, right=border_thin)
        
        for r_idx, row_val in enumerate(data_rows, 5):
            for c_idx, val in enumerate(row_val, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.font = Font(name="Segoe UI", size=10)
                cell.border = data_border
                
                # Column Specific Alignments & Format
                if c_idx in [4, 5]: # Match Score, Deal Value
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx in [6, 7]: # Priority, Status
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    
                # Format Percentage
                if c_idx == 4:
                    if isinstance(val, (int, float)):
                        if val > 1:
                            cell.value = val / 100.0
                        cell.number_format = '0%'
                # Format Currency
                elif c_idx == 5:
                    if isinstance(val, (int, float)):
                        cell.number_format = '$#,##0'
            
            # Alternating row background fills
            if r_idx % 2 == 0:
                alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
                for col_idx in range(1, 9):
                    ws.cell(row=r_idx, column=col_idx).fill = alt_fill
                    
            ws.row_dimensions[r_idx].height = 20
            
        # Set explicit column widths based on longest value length
        for col in ws.columns:
            max_len = 0
            for cell in col:
                # Calculate max length only below the merged headers
                if cell.row > 3 and cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 11)
            
        wb.save(file_path)
